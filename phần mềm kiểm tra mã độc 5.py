import os
import hashlib
from docx import Document
import openpyxl
import PyPDF2

# Từ khóa nghi ngờ
suspicious_keywords = ["exec(", "eval(", "base64", "subprocess", "import os", "os.system", "compile(", "globals("]

# Hàm tạo mã băm SHA-256
def hash_file(file_path):
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(4096):
            sha256.update(chunk)
    return sha256.hexdigest()

# Hàm đọc nội dung file
def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext in [".txt", ".py", ".csv", ".log"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        elif ext == ".docx":
            doc = Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs)
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = ""
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    line = " ".join(str(cell) for cell in row if cell)
                    text += line + "\n"
            return text
        elif ext == ".pdf":
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                return "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            return ""
    except Exception as e:
        print(f" Lỗi đọc {file_path}: {e}")
        return ""

# Kiểm tra file và in ra màn hình
def check_file(file_path):
    content = extract_text(file_path)
    if not content:
        print(f"️ Không thể đọc nội dung file: {file_path}")
        return

    found = [kw for kw in suspicious_keywords if kw in content]
    file_hash = hash_file(file_path)

    print(f"\n File: {file_path}")
    print(f" Mã băm SHA-256: {file_hash}")
    if found:
        print(f" Phát hiện mã độc! Từ khóa: {', '.join(found)}")
    else:
        print(" Không phát hiện mã độc.")

# Quét toàn bộ thư mục
def scan_folder(folder_path):
    print(f"\n Bắt đầu quét thư mục: {folder_path}\n")
    for root, _, files in os.walk(folder_path):
        for name in files:
            path = os.path.join(root, name)
            check_file(path)
    print("\n Đã quét xong!\n")

# --- CHẠY ---
folder = input("Nhập đường dẫn thư mục cần kiểm tra (ví dụ: E:/Test): ")
scan_folder(folder)
